"""Provision the class accounts listed in an Excel workbook to Supabase Auth.

The script never prints email addresses, passwords, API keys, or API response
bodies. It defaults to a dry run; account creation requires --apply and a typed
project-ref confirmation.
"""

from __future__ import annotations

import argparse
import base64
import getpass
import json
import sys
import time
from collections import Counter
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from openpyxl import load_workbook


DEFAULT_PROJECT_REF = "ujhcwxscxepbttqcysal"
DEFAULT_WORKBOOK = Path.home() / "Downloads" / "6학년 학생 계정 목록.xlsx"
# Stay at the Auth Admin API's documented default page size for compatibility.
# Some project configurations reject larger `per_page` values.
AUTH_PAGE_SIZE = 50
PROFILE_LIMIT = 1000
PASSWORD_PREFIX = "2026"


class ProvisionError(Exception):
    pass


def normalized_header(value: Any) -> str:
    return str(value or "").strip().lower().replace(" ", "").replace("_", "")


def find_column(
    headers: list[Any],
    candidates: set[str],
    *,
    label: str,
    required: bool = True,
) -> int | None:
    normalized = [normalized_header(item) for item in headers]
    for candidate in candidates:
        key = normalized_header(candidate)
        if key in normalized:
            return normalized.index(key)
    if required:
        raise ProvisionError(f"엑셀에서 '{label}' 열을 찾지 못했습니다. 열 제목을 확인하세요.")
    return None


def int_cell(value: Any, label: str, row_number: int) -> int:
    try:
        number = int(value)
        if float(value) != number:
            raise ValueError
        return number
    except (TypeError, ValueError, OverflowError):
        raise ProvisionError(f"엑셀 {row_number}행의 {label} 값이 올바른 정수가 아닙니다.") from None


def load_students(workbook_path: Path) -> list[dict[str, Any]]:
    if not workbook_path.is_file():
        raise ProvisionError(f"엑셀 파일을 찾을 수 없습니다: {workbook_path}")

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = list(next(rows))
    except StopIteration:
        raise ProvisionError("엑셀 시트가 비어 있습니다.") from None

    class_col = find_column(
        headers, {"반", "학급", "class", "class_no", "classno"}, label="반"
    )
    number_col = find_column(
        headers, {"번호", "학생번호", "student_no", "studentno", "no"}, label="번호"
    )
    email_col = find_column(
        headers,
        {"이메일", "계정", "계정명", "아이디", "email", "emailaddress", "username"},
        label="이메일/계정명",
    )
    password_col = find_column(
        headers, {"비밀번호", "암호", "password", "pw"}, label="비밀번호"
    )
    name_col = find_column(
        headers,
        {"이름", "학생명", "student_name", "studentname", "name"},
        label="이름",
        required=False,
    )

    students: list[dict[str, Any]] = []
    seen_emails: set[str] = set()
    seen_rosters: set[tuple[int, int, int]] = set()
    for row_number, row in enumerate(rows, start=2):
        if not row or all(value is None or str(value).strip() == "" for value in row):
            continue
        class_no = int_cell(row[class_col], "반", row_number)  # type: ignore[index]
        student_no = int_cell(row[number_col], "번호", row_number)  # type: ignore[index]
        email = str(row[email_col] or "").strip().lower()  # type: ignore[index]
        source_password = str(row[password_col] or "")  # type: ignore[index]
        if class_no < 1 or class_no > 20 or student_no < 1 or student_no > 50:
            raise ProvisionError(f"엑셀 {row_number}행의 반 또는 번호가 허용 범위를 벗어났습니다.")
        if "@" not in email or email.startswith("@") or email.endswith("@") or not source_password:
            raise ProvisionError(f"엑셀 {row_number}행의 이메일 또는 비밀번호가 비어 있거나 올바르지 않습니다.")
        password = PASSWORD_PREFIX + source_password
        if len(password) != 8:
            raise ProvisionError(f"엑셀 {row_number}행의 비밀번호는 2026을 붙인 뒤 8자가 되어야 합니다.")
        roster_key = (6, class_no, student_no)
        if email in seen_emails or roster_key in seen_rosters:
            raise ProvisionError(f"엑셀 {row_number}행에서 이메일 또는 반·번호 중복을 찾았습니다.")
        seen_emails.add(email)
        seen_rosters.add(roster_key)

        # The supplied workbook has no student-name column. Use a non-identifying
        # roster label so the app's required student_profiles row can be created.
        raw_name = str(row[name_col] or "").strip() if name_col is not None else ""
        student_name = raw_name or f"6-{class_no}-{student_no:02d}"
        if not 2 <= len(student_name) <= 30:
            raise ProvisionError(f"엑셀 {row_number}행의 이름 또는 생성된 표시명이 2~30자 범위를 벗어났습니다.")

        students.append(
            {
                "grade": 6,
                "class_no": class_no,
                "student_no": student_no,
                "email": email,
                "password": password,
                "student_name": student_name,
                "row_number": row_number,
            }
        )

    if not students:
        raise ProvisionError("발급할 학생 행이 없습니다.")
    return students


def api_request(base_url: str, api_key: str, path: str, *, method: str = "GET", body: Any = None) -> Any:
    url = f"{base_url.rstrip('/')}{path}"
    payload = None if body is None else json.dumps(body, ensure_ascii=False).encode("utf-8")
    headers = {
        "apikey": api_key,
        "Content-Type": "application/json",
        "Accept": "application/json",
        "User-Agent": "student-account-provisioner/1.0",
    }
    # New sb_secret keys are opaque API keys, not JWTs. Supabase requires
    # them in apikey only; the Authorization bearer header is for JWT keys.
    if not api_key.startswith("sb_secret_"):
        headers["Authorization"] = f"Bearer {api_key}"
    request = Request(url, data=payload, method=method, headers=headers)
    try:
        with urlopen(request, timeout=30) as response:
            data = response.read()
            return json.loads(data.decode("utf-8")) if data else None
    except HTTPError as error:
        # Do not print the response body: it may contain account identifiers.
        raise ProvisionError(f"Supabase 요청 실패 (HTTP {error.code}, {method} {path.split('?')[0]}).") from None
    except (URLError, TimeoutError, OSError):
        raise ProvisionError(f"Supabase 연결에 실패했습니다 ({method} {path.split('?')[0]}).") from None


def list_auth_users(base_url: str, api_key: str) -> list[dict[str, Any]]:
    users: list[dict[str, Any]] = []
    page = 1
    while True:
        query = urlencode({"page": page, "per_page": AUTH_PAGE_SIZE})
        result = api_request(base_url, api_key, f"/auth/v1/admin/users?{query}")
        batch = result.get("users", []) if isinstance(result, dict) else []
        if not isinstance(batch, list):
            raise ProvisionError("Auth 사용자 목록 응답을 읽을 수 없습니다.")
        users.extend(batch)
        if len(batch) < AUTH_PAGE_SIZE:
            return users
        page += 1


def list_student_profiles(base_url: str, api_key: str) -> list[dict[str, Any]]:
    query = urlencode({"select": "user_id,grade,class_no,student_no", "limit": str(PROFILE_LIMIT)})
    profiles = api_request(base_url, api_key, f"/rest/v1/student_profiles?{query}")
    if not isinstance(profiles, list):
        raise ProvisionError("student_profiles 조회에 실패했습니다. 테이블과 관리자 키 설정을 확인하세요.")
    if len(profiles) == PROFILE_LIMIT:
        raise ProvisionError("학생 프로필이 1,000개 이상입니다. 스크립트의 페이지 조회 설정을 확인해야 합니다.")
    return profiles


def ensure_student_profile(base_url: str, api_key: str, user_id: str, student: dict[str, Any]) -> bool:
    query = urlencode(
        {
            "select": "user_id,grade,class_no,student_no",
            "user_id": f"eq.{user_id}",
        }
    )
    existing = api_request(base_url, api_key, f"/rest/v1/student_profiles?{query}")
    if not isinstance(existing, list):
        raise ProvisionError("학생 프로필 생성 여부를 확인할 수 없습니다.")
    expected = (6, student["class_no"], student["student_no"])
    if existing:
        profile = existing[0]
        actual = (profile.get("grade"), profile.get("class_no"), profile.get("student_no"))
        if actual != expected:
            raise ProvisionError("기존 학생 프로필의 반·번호가 엑셀과 달라 작업을 중단했습니다.")
        return False

    api_request(
        base_url,
        api_key,
        "/rest/v1/student_profiles",
        method="POST",
        body={
            "user_id": user_id,
            "student_name": student["student_name"],
            "grade": 6,
            "class_no": student["class_no"],
            "student_no": student["student_no"],
        },
    )
    return True


def main() -> int:
    parser = argparse.ArgumentParser(description="엑셀의 6학년 학생 계정을 Supabase에 일괄 발급합니다.")
    parser.add_argument("workbook", nargs="?", type=Path, default=DEFAULT_WORKBOOK, help="계정 목록 xlsx 파일 경로")
    parser.add_argument("--project-ref", default=DEFAULT_PROJECT_REF, help="Supabase project ref")
    parser.add_argument("--apply", action="store_true", help="실제 계정과 프로필을 생성합니다. 생략하면 사전 점검만 합니다.")
    args = parser.parse_args()

    creation_started = False
    try:
        students = load_students(args.workbook)
        base_url = f"https://{args.project_ref}.supabase.co"
        print(f"대상 프로젝트: {args.project_ref}")
        print(f"엑셀 행 수: {len(students)}")
        for class_no, count in sorted(Counter(student["class_no"] for student in students).items()):
            print(f"  {class_no}반: {count}명")
        print("표시 이름: 엑셀에 이름 열이 없으면 반·번호 형식으로 생성")

        api_key = getpass.getpass("Supabase secret/service_role 키를 입력하세요 (화면에 표시되지 않습니다): ").strip()
        if not api_key:
            raise ProvisionError("관리자 키가 비어 있습니다.")
        if api_key.startswith("sb_publishable_"):
            raise ProvisionError("publishable 키는 계정 발급 권한이 없습니다. secret 또는 service_role 키를 입력하세요.")
        if not api_key.startswith("sb_secret_"):
            try:
                payload_part = api_key.split(".")[1]
                claims = json.loads(base64.urlsafe_b64decode(payload_part + "=" * (-len(payload_part) % 4)))
            except (IndexError, ValueError, json.JSONDecodeError):
                raise ProvisionError("키 형식을 확인할 수 없습니다. 프로젝트의 secret 또는 service_role 키를 입력하세요.") from None
            if claims.get("role") != "service_role":
                raise ProvisionError("입력한 JWT 키에 service_role 권한이 없습니다.")

        auth_users = list_auth_users(base_url, api_key)
        profiles = list_student_profiles(base_url, api_key)
        users_by_email = {str(user.get("email", "")).strip().lower(): user for user in auth_users}
        profiles_by_user = {str(profile.get("user_id")): profile for profile in profiles}
        profiles_by_roster = {
            (profile.get("grade"), profile.get("class_no"), profile.get("student_no")): profile
            for profile in profiles
        }

        already_complete = 0
        repair_profiles = 0
        create_users = 0
        for student in students:
            roster = (6, student["class_no"], student["student_no"])
            auth_user = users_by_email.get(student["email"])
            roster_profile = profiles_by_roster.get(roster)
            if roster_profile and (not auth_user or str(roster_profile.get("user_id")) != str(auth_user.get("id"))):
                raise ProvisionError("엑셀의 반·번호가 다른 기존 Auth 계정의 프로필과 겹쳐 작업을 중단했습니다.")
            if auth_user:
                profile = profiles_by_user.get(str(auth_user.get("id")))
                if profile:
                    actual = (profile.get("grade"), profile.get("class_no"), profile.get("student_no"))
                    if actual != roster:
                        raise ProvisionError("기존 Auth 계정의 프로필이 엑셀의 반·번호와 달라 작업을 중단했습니다.")
                    already_complete += 1
                else:
                    repair_profiles += 1
            else:
                create_users += 1

        print(f"이미 계정과 프로필이 연결됨: {already_complete}명")
        print(f"기존 계정에 프로필만 연결: {repair_profiles}명")
        print(f"새 계정 생성: {create_users}명")
        if not args.apply:
            print("사전 점검 완료. 실제 발급은 --apply 옵션으로 실행하세요.")
            return 0

        confirmation = input(f"실제 발급을 시작하려면 프로젝트 ref '{args.project_ref}'를 그대로 입력하세요: ").strip()
        if confirmation != args.project_ref:
            raise ProvisionError("프로젝트 확인 문구가 일치하지 않아 생성하지 않았습니다.")

        creation_started = True
        created = 0
        repaired = 0
        skipped = 0
        for index, student in enumerate(students, start=1):
            auth_user = users_by_email.get(student["email"])
            if auth_user:
                profile = profiles_by_user.get(str(auth_user.get("id")))
                if profile:
                    skipped += 1
                    continue
                user_id = str(auth_user.get("id"))
                if ensure_student_profile(base_url, api_key, user_id, student):
                    repaired += 1
                continue

            result = api_request(
                base_url,
                api_key,
                "/auth/v1/admin/users",
                method="POST",
                body={
                    "email": student["email"],
                    "password": student["password"],
                    "email_confirm": True,
                    "user_metadata": {
                        "student_name": student["student_name"],
                        "grade": 6,
                        "class_no": student["class_no"],
                        "student_no": student["student_no"],
                    },
                },
            )
            user = result.get("user", result) if isinstance(result, dict) else None
            user_id = str(user.get("id", "")) if isinstance(user, dict) else ""
            if not user_id:
                raise ProvisionError(f"{index}/{len(students)}번째 계정 생성 응답에 사용자 ID가 없어 중단했습니다.")
            ensure_student_profile(base_url, api_key, user_id, student)
            created += 1
            # Keep the burst below common Auth rate limits.
            if index < len(students):
                time.sleep(0.25)

        print(f"완료: 새 계정 {created}명, 프로필 보완 {repaired}명, 이미 완료되어 건너뜀 {skipped}명")
        return 0
    except ProvisionError as error:
        print(f"오류: {error}", file=sys.stderr)
        if creation_started:
            print("생성 도중 오류가 났다면 일부 항목이 처리되었을 수 있습니다. 다시 실행하면 완료된 계정은 건너뜁니다.", file=sys.stderr)
        return 1
    except Exception:
        # Do not emit tracebacks that could include the in-memory account data.
        print("예상하지 못한 오류가 발생했습니다. 민감한 값은 표시하지 않았습니다.", file=sys.stderr)
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
